#!/usr/bin/env python3
"""
Legal Document Generator Web Application
Integrates with Google Gemini and Claude Sonnet APIs for intelligent document generation
"""

import os
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional, Any
import hashlib

# Web framework and utilities
from flask import Flask, render_template, request, jsonify, flash, redirect, url_for, send_file
from werkzeug.utils import secure_filename
from werkzeug.exceptions import RequestEntityTooLarge

# Document processing
import PyPDF2
from PIL import Image
import pytesseract
from docx import Document as DocxDocument
from docx.shared import Inches

# AI/ML libraries
import google.generativeai as genai
from anthropic import Anthropic
import tiktoken

# Database
import sqlite3
from contextlib import contextmanager

# Configuration
from dataclasses import dataclass
from enum import Enum

# Logging setup
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Configuration
@dataclass
class Config:
    SECRET_KEY: str = os.getenv('SECRET_KEY', 'your-secret-key-here')
    GEMINI_API_KEY: str = os.getenv('GEMINI_API_KEY', '')
    CLAUDE_API_KEY: str = os.getenv('CLAUDE_API_KEY', '')
    UPLOAD_FOLDER: str = 'uploads'
    OUTPUT_FOLDER: str = 'generated_docs'
    MAX_CONTENT_LENGTH: int = 16 * 1024 * 1024  # 16MB
    ALLOWED_EXTENSIONS: set = {'pdf', 'docx', 'txt', 'png', 'jpg', 'jpeg'}
    DATABASE_PATH: str = 'legal_docs.db'

class LLMProvider(Enum):
    GEMINI = "gemini"
    CLAUDE = "claude"

class DocumentType(Enum):
    CONTRACT = "contract"
    AGREEMENT = "agreement"
    LEGAL_BRIEF = "legal_brief"
    MOTION = "motion"
    LEASE = "lease"
    NDA = "nda"
    TERMS_OF_SERVICE = "terms_of_service"

# Initialize Flask app
app = Flask(__name__)
config = Config()
app.config['SECRET_KEY'] = config.SECRET_KEY
app.config['MAX_CONTENT_LENGTH'] = config.MAX_CONTENT_LENGTH

# Create necessary directories
Path(config.UPLOAD_FOLDER).mkdir(exist_ok=True)
Path(config.OUTPUT_FOLDER).mkdir(exist_ok=True)

# Initialize AI clients
if config.GEMINI_API_KEY:
    genai.configure(api_key=config.GEMINI_API_KEY)
    gemini_model = genai.GenerativeModel('gemini-pro')
else:
    gemini_model = None
    logger.warning("Gemini API key not provided")

if config.CLAUDE_API_KEY:
    claude_client = Anthropic(api_key=config.CLAUDE_API_KEY)
else:
    claude_client = None
    logger.warning("Claude API key not provided")

# Database setup
def init_database():
    """Initialize the SQLite database with required tables"""
    with sqlite3.connect(config.DATABASE_PATH) as conn:
        cursor = conn.cursor()
        
        # Documents table
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                filename TEXT NOT NULL,
                original_filename TEXT NOT NULL,
                file_path TEXT NOT NULL,
                file_type TEXT NOT NULL,
                document_type TEXT,
                content_text TEXT,
                metadata TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                file_hash TEXT UNIQUE
            )
        ''')
        
        # Generated documents table
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS generated_documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                content TEXT NOT NULL,
                document_type TEXT NOT NULL,
                llm_provider TEXT NOT NULL,
                source_documents TEXT,
                parameters TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                file_path TEXT
            )
        ''')
        
        conn.commit()

@contextmanager
def get_db_connection():
    """Context manager for database connections"""
    conn = sqlite3.connect(config.DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()

class DocumentProcessor:
    """Handles document processing and text extraction"""
    
    @staticmethod
    def allowed_file(filename: str) -> bool:
        return '.' in filename and \
               filename.rsplit('.', 1)[1].lower() in config.ALLOWED_EXTENSIONS
    
    @staticmethod
    def get_file_hash(file_path: str) -> str:
        """Generate MD5 hash of file for duplicate detection"""
        hash_md5 = hashlib.md5()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(4096), b""):
                hash_md5.update(chunk)
        return hash_md5.hexdigest()
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text from PDF files"""
        try:
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text.strip()
        except Exception as e:
            logger.error(f"Error extracting text from PDF: {e}")
            return ""
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX files"""
        try:
            doc = DocxDocument(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {e}")
            return ""
    
    @staticmethod
    def extract_text_from_image(file_path: str) -> str:
        """Extract text from images using OCR"""
        try:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting text from image: {e}")
            return ""
    
    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """Extract text from plain text files"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read().strip()
        except Exception as e:
            logger.error(f"Error reading text file: {e}")
            return ""
    
    def process_document(self, file_path: str, file_type: str) -> str:
        """Process document and extract text based on file type"""
        if file_type == 'pdf':
            return self.extract_text_from_pdf(file_path)
        elif file_type == 'docx':
            return self.extract_text_from_docx(file_path)
        elif file_type in ['png', 'jpg', 'jpeg']:
            return self.extract_text_from_image(file_path)
        elif file_type == 'txt':
            return self.extract_text_from_txt(file_path)
        else:
            return ""

class LLMService:
    """Service for interacting with different LLM providers"""
    
    def __init__(self):
        self.processor = DocumentProcessor()
    
    def count_tokens(self, text: str) -> int:
        """Estimate token count for text"""
        try:
            encoding = tiktoken.get_encoding("cl100k_base")
            return len(encoding.encode(text))
        except:
            # Fallback: rough estimation
            return len(text.split()) * 1.3
    
    def generate_with_gemini(self, prompt: str, context_docs: List[str] = None) -> str:
        """Generate document using Google Gemini"""
        if not gemini_model:
            raise ValueError("Gemini API not configured")
        
        try:
            # Prepare context if documents are provided
            full_prompt = prompt
            if context_docs:
                context = "\n\n".join([f"Document {i+1}:\n{doc}" for i, doc in enumerate(context_docs)])
                full_prompt = f"Context Documents:\n{context}\n\nTask: {prompt}"
            
            response = gemini_model.generate_content(full_prompt)
            return response.text
        except Exception as e:
            logger.error(f"Error with Gemini API: {e}")
            raise
    
    def generate_with_claude(self, prompt: str, context_docs: List[str] = None) -> str:
        """Generate document using Claude Sonnet"""
        if not claude_client:
            raise ValueError("Claude API not configured")
        
        try:
            # Prepare context if documents are provided
            full_prompt = prompt
            if context_docs:
                context = "\n\n".join([f"Document {i+1}:\n{doc}" for i, doc in enumerate(context_docs)])
                full_prompt = f"Context Documents:\n{context}\n\nTask: {prompt}"
            
            response = claude_client.messages.create(
                model="claude-3-sonnet-20240229",
                max_tokens=4000,
                messages=[
                    {"role": "user", "content": full_prompt}
                ]
            )
            return response.content[0].text
        except Exception as e:
            logger.error(f"Error with Claude API: {e}")
            raise
    
    def generate_document(self, provider: LLMProvider, prompt: str, 
                         context_docs: List[str] = None) -> str:
        """Generate document using specified LLM provider"""
        if provider == LLMProvider.GEMINI:
            return self.generate_with_gemini(prompt, context_docs)
        elif provider == LLMProvider.CLAUDE:
            return self.generate_with_claude(prompt, context_docs)
        else:
            raise ValueError(f"Unsupported LLM provider: {provider}")

# Initialize services
llm_service = LLMService()

# Routes
@app.route('/')
def index():
    """Main dashboard"""
    with get_db_connection() as conn:
        cursor = conn.cursor()
        
        # Get document counts
        cursor.execute("SELECT COUNT(*) as count FROM documents")
        doc_count = cursor.fetchone()['count']
        
        cursor.execute("SELECT COUNT(*) as count FROM generated_documents")
        generated_count = cursor.fetchone()['count']
        
        # Get recent documents
        cursor.execute("""
            SELECT * FROM documents 
            ORDER BY created_at DESC 
            LIMIT 5
        """)
        recent_docs = cursor.fetchall()
        
        cursor.execute("""
            SELECT * FROM generated_documents 
            ORDER BY created_at DESC 
            LIMIT 5
        """)
        recent_generated = cursor.fetchall()
    
    return render_template('dashboard.html', 
                         doc_count=doc_count, 
                         generated_count=generated_count,
                         recent_docs=recent_docs,
                         recent_generated=recent_generated)

@app.route('/upload', methods=['GET', 'POST'])
def upload_document():
    """Upload and process legal documents"""
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('No file selected', 'error')
            return redirect(request.url)
        
        file = request.files['file']
        if file.filename == '':
            flash('No file selected', 'error')
            return redirect(request.url)
        
        if file and DocumentProcessor.allowed_file(file.filename):
            try:
                # Save file
                filename = secure_filename(file.filename)
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"{timestamp}_{filename}"
                file_path = os.path.join(config.UPLOAD_FOLDER, filename)
                file.save(file_path)
                
                # Process document
                processor = DocumentProcessor()
                file_type = filename.rsplit('.', 1)[1].lower()
                content_text = processor.process_document(file_path, file_type)
                file_hash = processor.get_file_hash(file_path)
                
                # Get document type from form
                document_type = request.form.get('document_type', '')
                
                # Store in database
                with get_db_connection() as conn:
                    cursor = conn.cursor()
                    try:
                        cursor.execute("""
                            INSERT INTO documents 
                            (filename, original_filename, file_path, file_type, 
                             document_type, content_text, file_hash)
                            VALUES (?, ?, ?, ?, ?, ?, ?)
                        """, (filename, file.filename, file_path, file_type, 
                              document_type, content_text, file_hash))
                        conn.commit()
                        flash('Document uploaded and processed successfully!', 'success')
                    except sqlite3.IntegrityError:
                        flash('This document has already been uploaded.', 'warning')
                        os.remove(file_path)  # Remove duplicate file
                
                return redirect(url_for('view_documents'))
                
            except Exception as e:
                logger.error(f"Error processing upload: {e}")
                flash(f'Error processing document: {str(e)}', 'error')
        else:
            flash('Invalid file type. Allowed types: PDF, DOCX, TXT, PNG, JPG, JPEG', 'error')
    
    return render_template('upload.html', document_types=DocumentType)

@app.route('/documents')
def view_documents():
    """View all uploaded documents"""
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT * FROM documents 
            ORDER BY created_at DESC
        """)
        documents = cursor.fetchall()
    
    return render_template('documents.html', documents=documents)

@app.route('/generate', methods=['GET', 'POST'])
def generate_document():
    """Generate new legal documents"""
    if request.method == 'POST':
        try:
            # Get form data
            title = request.form.get('title')
            document_type = request.form.get('document_type')
            llm_provider = LLMProvider(request.form.get('llm_provider'))
            custom_prompt = request.form.get('custom_prompt')
            source_doc_ids = request.form.getlist('source_documents')
            
            # Get source documents
            context_docs = []
            if source_doc_ids:
                with get_db_connection() as conn:
                    cursor = conn.cursor()
                    placeholders = ','.join(['?' for _ in source_doc_ids])
                    cursor.execute(f"""
                        SELECT content_text FROM documents 
                        WHERE id IN ({placeholders})
                    """, source_doc_ids)
                    context_docs = [row['content_text'] for row in cursor.fetchall()]
            
            # Build prompt
            base_prompt = f"""
            Generate a professional {document_type} document with the title "{title}".
            
            Requirements:
            - Follow standard legal document formatting
            - Include all necessary legal clauses and provisions
            - Ensure professional language and structure
            - Make it comprehensive and legally sound
            
            Additional Instructions: {custom_prompt}
            
            Please generate a complete, professional document.
            """
            
            # Generate document
            generated_content = llm_service.generate_document(
                provider=llm_provider,
                prompt=base_prompt,
                context_docs=context_docs if context_docs else None
            )
            
            # Save generated document to file
            output_filename = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{secure_filename(title)}.txt"
            output_path = os.path.join(config.OUTPUT_FOLDER, output_filename)
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(generated_content)
            
            # Store in database
            with get_db_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO generated_documents 
                    (title, content, document_type, llm_provider, source_documents, 
                     parameters, file_path)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (title, generated_content, document_type, llm_provider.value,
                      ','.join(source_doc_ids), custom_prompt, output_path))
                conn.commit()
            
            flash('Document generated successfully!', 'success')
            return redirect(url_for('view_generated'))
            
        except Exception as e:
            logger.error(f"Error generating document: {e}")
            flash(f'Error generating document: {str(e)}', 'error')
    
    # Get available source documents
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, original_filename, document_type FROM documents")
        source_documents = cursor.fetchall()
    
    return render_template('generate.html', 
                         document_types=DocumentType,
                         llm_providers=LLMProvider,
                         source_documents=source_documents)

@app.route('/generated')
def view_generated():
    """View all generated documents"""
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT * FROM generated_documents 
            ORDER BY created_at DESC
        """)
        documents = cursor.fetchall()
    
    return render_template('generated.html', documents=documents)

@app.route('/document/<int:doc_id>')
def view_document_detail(doc_id):
    """View detailed information about a specific document"""
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM documents WHERE id = ?", (doc_id,))
        document = cursor.fetchone()
        
        if not document:
            flash('Document not found', 'error')
            return redirect(url_for('view_documents'))
    
    return render_template('document_detail.html', document=document)

@app.route('/generated/<int:doc_id>')
def view_generated_detail(doc_id):
    """View detailed information about a specific generated document"""
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM generated_documents WHERE id = ?", (doc_id,))
        document = cursor.fetchone()
        
        if not document:
            flash('Document not found', 'error')
            return redirect(url_for('view_generated'))
    
    return render_template('generated_detail.html', document=document)

@app.route('/download/<int:doc_id>')
def download_generated(doc_id):
    """Download a generated document"""
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM generated_documents WHERE id = ?", (doc_id,))
        document = cursor.fetchone()
        
        if not document or not os.path.exists(document['file_path']):
            flash('Document not found', 'error')
            return redirect(url_for('view_generated'))
    
    return send_file(document['file_path'], as_attachment=True, 
                    download_name=f"{document['title']}.txt")

@app.route('/api/documents', methods=['GET'])
def api_get_documents():
    """API endpoint to get documents"""
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM documents ORDER BY created_at DESC")
        documents = [dict(row) for row in cursor.fetchall()]
    
    return jsonify(documents)

@app.route('/api/generate', methods=['POST'])
def api_generate_document():
    """API endpoint for document generation"""
    try:
        data = request.get_json()
        
        # Validate required fields
        required_fields = ['title', 'document_type', 'llm_provider']
        for field in required_fields:
            if field not in data:
                return jsonify({'error': f'Missing required field: {field}'}), 400
        
        # Generate document
        llm_provider = LLMProvider(data['llm_provider'])
        generated_content = llm_service.generate_document(
            provider=llm_provider,
            prompt=data.get('prompt', ''),
            context_docs=data.get('context_docs', [])
        )
        
        return jsonify({
            'success': True,
            'content': generated_content,
            'token_count': llm_service.count_tokens(generated_content)
        })
        
    except Exception as e:
        logger.error(f"API generation error: {e}")
        return jsonify({'error': str(e)}), 500

@app.errorhandler(413)
def too_large(e):
    flash('File is too large. Maximum size is 16MB.', 'error')
    return redirect(request.url), 413

@app.errorhandler(404)
def not_found(e):
    return render_template('404.html'), 404

@app.errorhandler(500)
def server_error(e):
    logger.error(f"Server error: {e}")
    return render_template('500.html'), 500

if __name__ == '__main__':
    # Initialize database
    init_database()
    
    # Check API configurations
    if not config.GEMINI_API_KEY and not config.CLAUDE_API_KEY:
        logger.warning("No API keys configured. Please set GEMINI_API_KEY and/or CLAUDE_API_KEY environment variables.")
    
    # Run the application
    app.run(debug=True, host='0.0.0.0', port=5000)
